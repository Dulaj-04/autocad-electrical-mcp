import asyncio

import openpyxl
import pytest

from acad_electrical_mcp import plan as plans
from acad_electrical_mcp import server
from acad_electrical_mcp.live import LiveSession, PlainCom
from acad_electrical_mcp.model import Store
from fake_acad import App


@pytest.fixture()
def s(tmp_path, monkeypatch):
    ws = tmp_path / "ws"
    monkeypatch.setattr(server.CFG, "workspace", ws)
    monkeypatch.setattr(server.CFG, "require_plan", False)
    monkeypatch.setattr(server, "STORE", Store(ws))
    monkeypatch.setattr(server, "PLANS", plans.PlanStore(ws))
    return server


def define_lighting(s):
    return s.plan_define("demo", "Floor 4 lighting", "Import the DIALux layout for 4F and total the load.",
                         ["4F luminaires placed and validated", "Lighting schedule produced"],
                         template="lighting_from_dialux")


def test_templates_expose_inputs_and_what_is_not_built(s):
    t = {x["template"]: x for x in s.plan_templates()["templates"]}
    assert {"lighting_from_dialux", "load_and_max_demand", "full_project"} <= set(t)
    assert t["cable_and_protection"]["not_built"]  # honest about gaps
    assert {w["kind"] for w in s.plan_templates()["workbooks"]} == {
        "load_schedule", "lighting_requirements", "luminaire_list", "cable_protection_inputs"}


def test_define_requires_goal_criteria_and_is_not_overwritten(s):
    with pytest.raises(ValueError, match="success criterion"):
        s.plan_define("demo", "t", "g", [])
    with pytest.raises(ValueError, match="Unknown template"):
        s.plan_define("demo", "t", "g", ["c"], template="nope")
    out = define_lighting(s)
    assert out["status"] == "draft" and out["approved"] is False
    assert out["progress"] == {"done": 0, "total": 9}
    with pytest.raises(ValueError, match="already exists"):
        define_lighting(s)
    assert s.plan_define("demo", "t", "g", ["c"], replace=True)["title"] == "t"


def test_plan_asks_the_user_for_exactly_what_is_missing(s):
    define_lighting(s)
    nxt = s.plan_next("demo")
    ids = {i["id"] for i in nxt["needed_from_user"]}
    assert {"arch_dxf", "dialux_export", "luminaire_list", "alignment_points"} <= ids
    lum = next(i for i in nxt["needed_from_user"] if i["id"] == "luminaire_list")
    assert lum["workbook_template"] == "luminaire_list" and "wattage" in lum["why"].lower()
    assert "Needed from you" in nxt["checklist"]


def test_steps_wait_for_approval_inputs_and_dependencies(s):
    define_lighting(s)
    with pytest.raises(ValueError, match="not approved"):
        s.plan_update_step("demo", "s2", "in_progress")
    with pytest.raises(ValueError, match="user_confirmed"):
        s.plan_approve("demo", user_confirmed=False)
    s.plan_approve("demo", user_confirmed=True, note="ok by user")
    with pytest.raises(ValueError, match="blocked"):
        s.plan_update_step("demo", "s2", "in_progress")  # dialux_export missing
    s.plan_provide_input("demo", "dialux_export", "C:/x/floor4.dxf")
    s.plan_update_step("demo", "s2", "in_progress")
    with pytest.raises(ValueError, match="result"):
        s.plan_update_step("demo", "s2", "done")
    out = s.plan_update_step("demo", "s2", "done", result="3 blocks, 68 luminaires")
    assert out["progress"]["done"] == 1
    with pytest.raises(ValueError, match="blocked"):
        s.plan_update_step("demo", "s3", "done", result="x")  # alignment points still missing
    s.plan_provide_input("demo", "alignment_points", "A(0,0)->(0,0); B(6,0)->(6000,0)")
    assert s.plan_update_step("demo", "s3", "done", result="residual 0")["progress"]["done"] == 2
    forced = s.plan_update_step("demo", "s9", "done", result="skipped check", note="user override",
                                force=True)
    assert next(x for x in forced["steps"] if x["id"] == "s9")["status"] == "done"


def test_inputs_notes_and_extra_steps(s):
    define_lighting(s)
    with pytest.raises(ValueError, match="Say what"):
        s.plan_provide_input("demo", "arch_dxf")
    s.plan_provide_input("demo", "lighting_requirements", not_applicable=True, note="not needed")
    s.plan_add_input("demo", "client_brief", "Client brief", "Sets scope", "Send the PDF", kind="file")
    with pytest.raises(ValueError, match="already exists"):
        s.plan_add_input("demo", "client_brief", "x", "y", "z")
    with pytest.raises(ValueError, match="Unknown input"):
        s.plan_add_step("demo", "Review brief", "review", needs_inputs=["nope"])
    added = s.plan_add_step("demo", "Review brief", "review", needs_inputs=["client_brief"])
    assert added["added"]["id"] == "x1"
    s.plan_note("demo", "Mounting height 2.7 m", "assumption", source="client call")
    shown = s.plan_show("demo")
    assert shown["assumptions"][0]["source"] == "client call"
    assert any(i["id"] == "client_brief" for i in shown["next"]["needed_from_user"])


def test_plan_is_persisted_per_project_and_validates_names(s):
    define_lighting(s)
    assert s.plan_show("demo")["title"] == "Floor 4 lighting"
    with pytest.raises(ValueError, match="No plan"):
        s.plan_show("other")
    with pytest.raises(ValueError, match="Invalid project"):
        s.plan_show("../x")


def test_prompts_are_registered(s):
    names = {p.name for p in asyncio.run(server.mcp.list_prompts())}
    assert {"plan_a_request", "lighting_from_dialux", "load_schedule_and_demand",
            "project_roadmap"} <= names
    assert "Request: wire the 4th floor" in server.plan_a_request("demo", "wire the 4th floor")
    assert "plan-first" in server.INSTRUCTIONS.lower()


# ------------------------------------------------------------------ Excel
def fill(path, rows, sheet="Data"):
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    for r, values in enumerate(rows, start=2):
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v)
    wb.save(path)


LOAD_OK = [
    ["4F", "DB-4F", "4F-LTG", "Lighting", "lighting", 1, 10, None, 0.8, "1", "N", "DIALux schedule"],
    ["GF", "MSB", "PUMP-1", "Pump 1", "pump", 1, 100, 0.85, 1.0, "3", "Y", "brief"],
    ["GF", "MSB", "PUMP-2", "Pump 2", "pump", 1, 100, 0.85, 1.0, "3", "Y", "brief"],
    ["4F", "DB-4F", "4F-AC", "AC units", "ac", 4, 5, 0.9, None, "3", "N", "drawing: 4 units"],
]


def test_workbook_template_is_created_once_with_instructions_and_formulas(s, tmp_path):
    r = s.workbook_create("demo", "load_schedule")
    wb = openpyxl.load_workbook(r["path"])
    assert wb.sheetnames == ["Instructions", "Data"]
    assert wb["Data"]["M2"].value.startswith("=IF(") and "Floor *" in wb["Data"]["A1"].value
    assert "Demand factor" in [c.value for c in wb["Data"][1]][8]
    with pytest.raises(ValueError, match="already exists"):
        s.workbook_create("demo", "load_schedule")
    assert s.workbook_create("demo", "load_schedule", overwrite=True)["path"] == r["path"]


def test_empty_template_is_not_ok_and_filled_one_is(s):
    r = s.workbook_create("demo", "load_schedule")
    assert s.workbook_read("demo", "load_schedule")["ok"] is False  # no rows yet
    fill(r["path"], LOAD_OK)
    res = s.workbook_read("demo", "load_schedule")
    assert res["ok"] and res["row_count"] == 4 and res["errors"] == 0
    assert res["rows"][3]["df"] is None  # blank stays blank, never defaulted


def test_validation_reports_row_and_column(s):
    r = s.workbook_create("demo", "load_schedule")
    fill(r["path"], [
        ["4F", "DB", "A1", "x", "lighting", 1, 10, None, 1.5, "1", "N", ""],      # df > 1
        ["4F", "DB", "A1", "x", "heater", 1, 10, None, 0.5, "1", "N", ""],         # dup id + bad category
        [None, "DB", "A3", "x", "ac", "two", 10, None, 0.5, "1", "N", ""],         # no floor, qty text
    ])
    res = s.workbook_read("demo", "load_schedule")
    msgs = [(i["row"], i["column"], i["message"]) for i in res["issues"]]
    assert not res["ok"]
    assert any(m[0] == 2 and "at most 1" in m[2] for m in msgs)
    assert any(m[0] == 3 and "not one of" in m[2] for m in msgs)
    assert any("Duplicate id 'a1'" in m[2] for m in msgs)
    assert any(m[0] == 4 and m[1].startswith("Floor") and "required" in m[2] for m in msgs)
    assert any(m[0] == 4 and "not a number" in m[2] for m in msgs)
    with pytest.raises(ValueError, match="error"):
        s.load_summary("demo")


def test_formula_without_stored_value_is_flagged(s):
    r = s.workbook_create("demo", "load_schedule")
    fill(r["path"], [["4F", "DB", "A1", "x", "ac", "=2*3", 10, None, 0.5, "1", "N", ""]])
    res = s.workbook_read("demo", "load_schedule")
    assert any("Formula without a stored value" in i["message"] for i in res["issues"])


def test_load_summary_uses_only_supplied_factors_and_says_what_is_missing(s):
    r = s.workbook_create("demo", "load_schedule")
    fill(r["path"], LOAD_OK)
    out = s.load_summary("demo")
    assert out["total"]["connected_kw"] == 10 + 200 + 20
    assert out["total"]["demand_kw"] == 8 + 200  # AC has no factor: left out, not assumed 1
    assert out["rows_without_demand_factor"] == ["4F-AC"] and out["complete"] is False
    assert any("NOT included" in n for n in out["notes"])
    assert out["essential"]["connected_kw"] == 200 and out["essential"]["demand_kva"] == pytest.approx(235.294, abs=1e-3)
    assert out["by_floor"]["GF"]["connected_kw"] == 200 and out["by_category"]["pump"]["rows"] == 2
    assert out["rows_without_power_factor"] == ["4F-LTG"]
    # filling in the missing factor completes it
    wb = openpyxl.load_workbook(r["path"])
    wb["Data"]["I5"] = 0.7
    wb.save(r["path"])
    out2 = s.load_summary("demo")
    assert out2["complete"] and out2["total"]["demand_kw"] == pytest.approx(8 + 200 + 14)


def test_workbook_read_marks_the_plan_input_provided(s):
    s.plan_define("demo", "Loads", "Load schedule", ["done"], template="load_and_max_demand")
    r = s.workbook_create("demo", "load_schedule")
    fill(r["path"], LOAD_OK)
    res = s.workbook_read("demo", "load_schedule", plan_input_id="load_schedule")
    assert res["plan_input_marked_provided"] == "load_schedule"
    inputs = {i["id"]: i for i in s.plan_show("demo")["inputs"]}
    assert inputs["load_schedule"]["status"] == "provided" and inputs["demand_factors"]["status"] == "missing"


def test_cable_sheet_reports_missing_table_values_instead_of_filling_them(s):
    r = s.workbook_create("demo", "cable_protection_inputs")
    fill(r["path"], [["C1", "DB-4F", "AC", 20, None, None, None, None, 35, "clipped direct", "4mm2",
                      None, None, None, None, None, None, 25, "BS 7671 Table 4D5"]])
    res = s.workbook_read("demo", "cable_protection_inputs")
    assert res["ok"] and res["warnings"] >= 1
    assert any("Table values missing" in i["message"] for i in res["issues"])
    wb = openpyxl.load_workbook(r["path"])
    assert "L2*M2*N2*O2" in wb["Data"]["T2"].value  # Iz formula is visible in Excel


def test_luminaire_list_template_feeds_the_dialux_importer(s):
    from acad_electrical_mcp import dialux
    r = s.workbook_create("demo", "luminaire_list")
    fill(r["path"], [["LUM_A", 3, 18.5, 2100, ""]])
    got = dialux.read_luminaire_list(r["path"], sheet="Data")
    assert got["rows"][0] == {"type": "LUM_A", "qty": 3.0, "watts": 18.5, "lumens": 2100.0}


# ------------------------------------------------------------------ strict mode
def test_strict_mode_blocks_changes_until_goal_is_approved(s, monkeypatch, tmp_path):
    monkeypatch.setattr(server.CFG, "require_plan", True)
    app = App()
    monkeypatch.setattr(server, "_LIVE", LiveSession(PlainCom(app), tmp_path / "st"))
    with pytest.raises(ValueError, match="requires an approved plan"):
        s.live_place("4F", "socket", 10, 10)
    with pytest.raises(ValueError, match="requires an approved plan"):
        s.apply_changes("demo", "4F", "cs-x")
    assert s.live_connect()["drawing"] == "Drawing1.dwg"  # reading is always allowed
    s.plan_define("demo", "Sockets", "Place sockets", ["placed"])
    with pytest.raises(ValueError, match="requires an approved plan"):
        s.live_place("4F", "socket", 10, 10)  # defined but not approved
    s.plan_approve("demo", user_confirmed=True)
    assert s.live_place("4F", "socket", 10, 10)["created"]
    monkeypatch.setattr(server.CFG, "require_plan", False)


def test_dry_run_import_is_allowed_in_strict_mode(s, monkeypatch, tmp_path):
    from test_dialux import make_export
    monkeypatch.setattr(server.CFG, "require_plan", True)
    monkeypatch.setattr(server, "_LIVE", LiveSession(PlainCom(App()), tmp_path / "st"))
    export = make_export(str(tmp_path / "d.dxf"))
    assert s.live_import_luminaires("4F", export, block_names=["LUM_A"])["dry_run"]
    with pytest.raises(ValueError, match="requires an approved plan"):
        s.live_import_luminaires("4F", export, block_names=["LUM_A"], dry_run=False)
    monkeypatch.setattr(server.CFG, "require_plan", False)
