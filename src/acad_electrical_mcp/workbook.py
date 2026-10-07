"""Excel input sheets: templates to fill in, a validating reader, and the load summary.

Templates contain instructions, dropdowns and formulas so the arithmetic is visible in Excel. The
reader never trusts formula results (openpyxl-written formulas have no cached values until Excel
saves the file): it reads the INPUT columns and the calculations are redone here. Nothing is ever
filled in for you: blank demand factors, power factors or table values are reported as missing.
"""

from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from .backends.base import BackendError
from .config import safe_name

ROWS = 300  # formula rows pre-filled in each template


def _col(key, header, required=False, type="text", enum=None, note="", formula=None, aliases=()):
    return {"key": key, "header": header, "required": required, "type": type, "enum": enum,
            "note": note, "formula": formula, "aliases": list(aliases)}


CATEGORIES = ["lighting", "sockets", "ac", "pump", "lift", "fire", "other"]

WORKBOOKS: dict[str, dict[str, Any]] = {
    "load_schedule": {
        "title": "Load schedule",
        "purpose": "Every connected load, per floor and board. Feeds connected load and maximum demand.",
        "instructions": [
            "One row per load (or per group of identical loads).",
            "Take lighting from the lighting schedule (watts / 1000 = kW) and socket/AC counts from the drawing.",
            "Demand factor (0-1) must come from your approved method. Leave blank if unknown: it is reported as missing and never assumed.",
            "Power factor: leave blank if unknown (kVA is then not computed for that row).",
            "Essential = Y for loads that must run on the generator.",
            "Grey columns are formulas; do not type in them.",
        ],
        "columns": [
            _col("floor", "Floor", True, note="e.g. GF, 1F, 4F"),
            _col("board", "DB / Board", False, note="Board feeding the load, e.g. DB-4F"),
            _col("load_id", "Load ID", True, note="Unique, e.g. 4F-L01"),
            _col("description", "Description", False),
            _col("category", "Category", True, "enum", CATEGORIES),
            _col("qty", "Qty", True, "number"),
            _col("unit_kw", "Unit power (kW)", True, "number", note="Per item; 100 for each 100 kW pump"),
            _col("pf", "Power factor", False, "factor", note="0 < PF <= 1"),
            _col("df", "Demand factor", False, "factor", note="0 < DF <= 1, from your method"),
            _col("phases", "Phases", False, "enum", ["1", "3"]),
            _col("essential", "Essential (Y/N)", False, "enum", ["Y", "N"]),
            _col("source", "Source / notes", False, note="Where the numbers come from"),
            _col("connected", "Connected kW", formula="=IF(A{r}=\"\",\"\",F{r}*G{r})"),
            _col("demand", "Demand kW", formula="=IF(OR(A{r}=\"\",I{r}=\"\"),\"\",M{r}*I{r})"),
            _col("kva", "Demand kVA", formula="=IF(OR(N{r}=\"\",H{r}=\"\"),\"\",N{r}/H{r})"),
        ],
    },
    "lighting_requirements": {
        "title": "Lighting requirements",
        "purpose": "The design basis per room (target illuminance). Lux is calculated in DIALux, not here.",
        "instructions": ["One row per room.", "Name the standard or client requirement in Source."],
        "columns": [
            _col("room_id", "Room ID", True), _col("room_name", "Room name", True),
            _col("function", "Function", False), _col("target_lux", "Target lux (Em)", True, "number"),
            _col("source", "Source of requirement", True), _col("area_m2", "Area (m2)", False, "number"),
            _col("notes", "Notes", False),
        ],
    },
    "luminaire_list": {
        "title": "Luminaire list",
        "purpose": "Wattage per luminaire type for the DIALux import. Same layout as a DIALux parts list.",
        "instructions": ["Type must match the block name in the DIALux export (see dialux_inspect).",
                         "Power is the total luminaire power including driver, as in the datasheet."],
        "columns": [
            _col("type", "Luminaire", True, aliases=("type", "name", "article")),
            _col("qty", "Quantity", False, "number"), _col("watts", "Power [W]", True, "number"),
            _col("lumens", "Luminous flux [lm]", False, "number"), _col("notes", "Notes", False),
        ],
    },
    "cable_protection_inputs": {
        "title": "Cable and protection inputs",
        "purpose": "Per circuit inputs for cable sizing and protection checks. Table values come from your standard.",
        "instructions": [
            "Design current Ib: enter it, or leave blank and fill kW, power factor, voltage and phases.",
            "Tabulated current, correction factors and mV/A/m must be copied from the standard you use; name the table/clause in Source. They are never filled in for you.",
            "Check columns (grey) show Ib <= In <= Iz and the voltage drop; have results reviewed.",
        ],
        "columns": [
            _col("circuit", "Circuit ID", True), _col("from_board", "From (DB)", True),
            _col("load", "Load description", False), _col("ib", "Design current Ib (A)", False, "number"),
            _col("kw", "Load (kW)", False, "number"), _col("pf", "Power factor", False, "factor"),
            _col("voltage", "Voltage (V)", False, "number"), _col("phases", "Phases", False, "enum", ["1", "3"]),
            _col("length_m", "Length (m)", True, "number"), _col("method", "Installation method", True),
            _col("cable", "Cable type / size", False), _col("it", "Tabulated current It (A)", False, "number"),
            _col("ca", "Correction: ambient", False, "number"), _col("cg", "Correction: grouping", False, "number"),
            _col("ci", "Correction: insulation", False, "number"), _col("mv", "mV/A/m", False, "number"),
            _col("max_vd", "Max voltage drop (%)", False, "number"), _col("in_a", "Protective device In (A)", False, "number"),
            _col("source", "Source of table values (table / clause)", True),
            _col("iz", "Iz = It x Ca x Cg x Ci (A)", formula="=IF(OR(L{r}=\"\",M{r}=\"\",N{r}=\"\",O{r}=\"\"),\"\",L{r}*M{r}*N{r}*O{r})"),
            _col("vd", "Voltage drop (V)", formula="=IF(OR(P{r}=\"\",D{r}=\"\",I{r}=\"\"),\"\",P{r}*D{r}*I{r}/1000)"),
            _col("ok", "Ib <= In <= Iz ?", formula="=IF(OR(D{r}=\"\",R{r}=\"\",T{r}=\"\"),\"\",IF(AND(D{r}<=R{r},R{r}<=T{r}),\"OK\",\"CHECK\"))"),
        ],
    },
}


def list_workbooks() -> list[dict[str, Any]]:
    return [{"kind": k, "title": v["title"], "purpose": v["purpose"],
             "columns": [c["header"] for c in v["columns"] if not c["formula"]],
             "required": [c["header"] for c in v["columns"] if c["required"]]}
            for k, v in WORKBOOKS.items()]


def _openpyxl():
    try:
        import openpyxl
    except ImportError as exc:
        raise BackendError("Excel support needs openpyxl: pip install openpyxl") from exc
    return openpyxl


def inputs_dir(workspace: Path, project: str) -> Path:
    return workspace / safe_name(project, "project") / "inputs"


def create(workspace: Path, project: str, kind: str, overwrite: bool = False) -> dict[str, Any]:
    if kind not in WORKBOOKS:
        raise ValueError(f"Unknown workbook {kind!r}; options: {sorted(WORKBOOKS)}")
    openpyxl = _openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    spec = WORKBOOKS[kind]
    path = inputs_dir(workspace, project) / f"{kind}.xlsx"
    if path.exists() and not overwrite:
        raise ValueError(f"{path} already exists. It may contain your data: pass overwrite=true "
                         "only if you want to start again.")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    info = wb.active
    info.title = "Instructions"
    info["A1"] = spec["title"]
    info["A1"].font = Font(bold=True, size=14)
    info["A2"] = spec["purpose"]
    for i, line in enumerate(spec["instructions"], start=4):
        info.cell(row=i, column=1, value=f"{i - 3}. {line}")
    info["A" + str(len(spec["instructions"]) + 6)] = "Status: DRAFT input sheet. Values are only as good as their sources."
    info.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Data")
    head_fill = PatternFill("solid", fgColor="1F3864")
    calc_fill = PatternFill("solid", fgColor="D9D9D9")
    for c, col in enumerate(spec["columns"], start=1):
        cell = ws.cell(row=1, column=c, value=col["header"] + (" *" if col["required"] else ""))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill if not col["formula"] else PatternFill("solid", fgColor="595959")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(c)].width = max(14, min(34, len(col["header"]) + 4))
        if col["note"]:
            from openpyxl.comments import Comment
            cell.comment = Comment(col["note"], "acad-electrical-mcp")
        if col["formula"]:
            for r in range(2, ROWS + 2):
                ws.cell(row=r, column=c, value=col["formula"].format(r=r)).fill = calc_fill
        if col["type"] == "enum" and col["enum"]:
            dv = DataValidation(type="list", formula1='"' + ",".join(col["enum"]) + '"',
                                allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(c)}2:{get_column_letter(c)}{ROWS + 1}")
        if col["type"] == "factor":
            dv = DataValidation(type="decimal", operator="between", formula1="0.001", formula2="1",
                                allow_blank=True, error="Enter a number above 0 and up to 1")
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(c)}2:{get_column_letter(c)}{ROWS + 1}")
    ws.freeze_panes = "A2"
    wb.save(str(path))
    return {"path": str(path), "kind": kind, "sheet": "Data", "required_columns": [
        c["header"] for c in spec["columns"] if c["required"]],
        "tell_the_user": f"Open {path} in Excel, fill the 'Data' sheet (columns marked * are "
                         "required; grey columns are formulas), save it, and tell me. Then I check "
                         "it with workbook_read."}


def _norm(h: Any) -> str:
    return re.sub(r"[\s*]+", " ", str(h or "").strip().lower()).strip()


def _number(v: Any) -> float | None:
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip().replace(",", "."))
    except ValueError:
        return None


def read(path: str, kind: str) -> dict[str, Any]:
    """Read and validate a filled-in workbook. Returns rows (input columns only) and issues."""
    if kind not in WORKBOOKS:
        raise ValueError(f"Unknown workbook {kind!r}; options: {sorted(WORKBOOKS)}")
    p = Path(path)
    if not p.exists():
        raise BackendError(f"File not found: {path}")
    openpyxl = _openpyxl()
    wb_v = openpyxl.load_workbook(str(p), data_only=True)
    wb_f = openpyxl.load_workbook(str(p), data_only=False)
    name = "Data" if "Data" in wb_v.sheetnames else wb_v.sheetnames[0]
    ws_v, ws_f = wb_v[name], wb_f[name]
    spec = WORKBOOKS[kind]
    cols = [c for c in spec["columns"] if not c["formula"]]
    header_row = list(ws_v.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    idx: dict[str, int] = {}
    for c in cols:
        names = {_norm(c["header"]), *(_norm(a) for a in c["aliases"])}
        for i, h in enumerate(header_row):
            if _norm(h) in names or _norm(h).split(" [")[0] in names:
                idx[c["key"]] = i
                break
    issues: list[dict[str, Any]] = []
    for c in cols:
        if c["required"] and c["key"] not in idx:
            issues.append({"severity": "error", "row": 1, "column": c["header"],
                           "message": f"Required column '{c['header']}' not found in sheet '{name}'."})
    rows, seen = [], defaultdict(list)
    for r_i, (rv, rf) in enumerate(zip(ws_v.iter_rows(min_row=2, values_only=True),
                                       ws_f.iter_rows(min_row=2, values_only=True)), start=2):
        if not any(rv[idx[c["key"]]] not in (None, "") for c in cols if c["key"] in idx):
            continue
        row: dict[str, Any] = {"_row": r_i}
        for c in cols:
            if c["key"] not in idx:
                continue
            raw, formula = rv[idx[c["key"]]], rf[idx[c["key"]]]
            if raw is None and isinstance(formula, str) and formula.startswith("="):
                issues.append({"severity": "error", "row": r_i, "column": c["header"],
                               "message": "Formula without a stored value: open the file in Excel "
                                          "and save it, or type the value."})
                raw = None
            val: Any = raw
            if raw in (None, ""):
                val = None
                if c["required"]:
                    issues.append({"severity": "error", "row": r_i, "column": c["header"],
                                   "message": f"'{c['header']}' is required."})
            elif c["type"] in ("number", "factor"):
                val = _number(raw)
                if val is None:
                    issues.append({"severity": "error", "row": r_i, "column": c["header"],
                                   "message": f"'{raw}' is not a number."})
                elif c["type"] == "factor" and not (0 < val <= 1):
                    issues.append({"severity": "error", "row": r_i, "column": c["header"],
                                   "message": f"{val} must be above 0 and at most 1."})
                    val = None
                elif c["type"] == "number" and val < 0:
                    issues.append({"severity": "error", "row": r_i, "column": c["header"],
                                   "message": f"{val} must not be negative."})
            elif c["type"] == "enum":
                s = str(raw).strip()
                match = next((e for e in c["enum"] if e.lower() == s.lower()), None)
                if match is None:
                    issues.append({"severity": "error", "row": r_i, "column": c["header"],
                                   "message": f"'{s}' is not one of {c['enum']}."})
                val = match
            else:
                val = str(raw).strip()
            row[c["key"]] = val
        key = {"load_schedule": "load_id", "cable_protection_inputs": "circuit",
               "lighting_requirements": "room_id", "luminaire_list": "type"}[kind]
        if row.get(key):
            seen[str(row[key]).lower()].append(r_i)
        rows.append(row)
    for k, rs in seen.items():
        if len(rs) > 1:
            issues.append({"severity": "error", "row": rs[0], "column": None,
                           "message": f"Duplicate id '{k}' on rows {rs}."})
    if kind == "cable_protection_inputs":
        for row in rows:
            if row.get("ib") is None and not (row.get("kw") and row.get("voltage")):
                issues.append({"severity": "warning", "row": row["_row"], "column": "Ib",
                               "message": "No design current and no kW + voltage to derive it."})
            if row.get("it") is None or row.get("mv") is None:
                issues.append({"severity": "warning", "row": row["_row"], "column": "It / mV/A/m",
                               "message": "Table values missing: enter them from your standard."})
    errors = sum(i["severity"] == "error" for i in issues)
    return {"path": str(p.resolve()), "kind": kind, "sheet": name, "rows": rows,
            "row_count": len(rows), "errors": errors,
            "warnings": sum(i["severity"] == "warning" for i in issues), "issues": issues[:200],
            "ok": errors == 0 and bool(rows),
            "status": "DRAFT"}


def load_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Connected load and demand from load-schedule rows. Demand uses ONLY the factors supplied:
    rows without one are listed, not assumed to be 1."""
    def blank() -> dict[str, float]:
        return {"connected_kw": 0.0, "demand_kw": 0.0, "demand_kva": 0.0, "rows": 0}

    by_floor, by_board, by_cat = defaultdict(blank), defaultdict(blank), defaultdict(blank)
    total, essential = blank(), blank()
    no_df, no_pf = [], []
    for r in rows:
        if r.get("qty") is None or r.get("unit_kw") is None:
            continue
        conn = r["qty"] * r["unit_kw"]
        df, pf = r.get("df"), r.get("pf")
        dem = conn * df if df else 0.0
        kva = dem / pf if (df and pf) else 0.0
        if not df:
            no_df.append(r["load_id"])
        if df and not pf:
            no_pf.append(r["load_id"])
        for bucket in (by_floor[r["floor"]], by_board[r.get("board") or "(no board)"],
                       by_cat[r["category"]], total, *([essential] if r.get("essential") == "Y" else [])):
            bucket["connected_kw"] += conn
            bucket["demand_kw"] += dem
            bucket["demand_kva"] += kva
            bucket["rows"] += 1

    def fin(d):
        return {k: round(v, 3) for k, v in d.items()}

    complete = not no_df
    return {
        "total": fin(total), "essential": fin(essential),
        "by_floor": {k: fin(v) for k, v in sorted(by_floor.items())},
        "by_board": {k: fin(v) for k, v in sorted(by_board.items())},
        "by_category": {k: fin(v) for k, v in sorted(by_cat.items())},
        "rows_without_demand_factor": no_df, "rows_without_power_factor": no_pf,
        "complete": complete,
        "notes": ["Maximum demand here = sum of (connected kW x the demand factors you supplied). "
                  "No diversity between boards is applied unless it is in your factors.",
                  "demand_kva uses the power factor on each row; rows without one add 0 kVA, so the "
                  "kVA total is understated until they are filled in."] + ([
                      f"{len(no_df)} row(s) have no demand factor: their demand is NOT included, so the "
                      "demand total is incomplete."] if no_df else []),
        "status": "DRAFT"}
